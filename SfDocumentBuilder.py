from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os


# ─────────────────────────────────────────────
#  Color palette  (hex strings, no leading #)
# ─────────────────────────────────────────────
COLORS = {
    # section accent colors
    "title_bg":      "1F3864",   # deep navy     — title bar
    "title_fg":      "FFFFFF",   # white
    "h1_bg":         "2E75B6",   # Salesforce blue  — Summary / Data Access Flow
    "h1_fg":         "FFFFFF",
    "h2_bg":         "D6E4F0",   # light blue   — section headings
    "h2_fg":         "1F3864",
    "h3_best":       "E2EFDA",   # soft green   — Best Practices
    "h3_best_fg":    "375623",
    "h3_limit":      "FCE4D6",   # soft orange  — Limitations
    "h3_limit_fg":   "833C00",
    "bullet_text":   "1F3864",   # navy body text
    "flow_bg":       "EBF3FB",   # very light blue — flow box
    "flow_border":   "2E75B6",
    "example_bg":    "FFF2CC",   # light yellow  — example box
    "example_fg":    "7F6000",
    "separator":     "2E75B6",
}


def _rgb(hex_str: str) -> RGBColor:
    r, g, b = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    return RGBColor(r, g, b)


def _set_cell_bg(cell, hex_color: str):
    """Apply a solid background shading to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_border(cell, hex_color: str, size: int = 12):
    """Draw a single border on all four sides of a cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), hex_color)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def _banner_paragraph(doc: Document, text: str,
                       bg: str, fg: str,
                       font_size: int = 13,
                       bold: bool = True) -> None:
    """
    One-cell table that looks like a coloured heading banner.
    Safer than paragraph shading which Word renders inconsistently.
    """
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _set_cell_bg(cell, bg)
    _set_cell_border(cell, bg, size=4)

    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(font_size)
    run.font.color.rgb = _rgb(fg)
    run.font.name = "Calibri"

    # Tight internal padding
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), "100")
        m.set(qn("w:type"), "dxa")
        tc_mar.append(m)
    tc_pr.append(tc_mar)

    doc.add_paragraph()          # breathing room after banner


def _bullet(doc: Document, text: str, fg_hex: str = "1F3864",
            indent_level: int = 0) -> None:
    """
    Add a coloured bullet paragraph using list-bullet style.
    """
    para = doc.add_paragraph(style="List Bullet")
    if indent_level:
        para.paragraph_format.left_indent = Inches(0.25 * indent_level)
    run = para.add_run(text)
    run.font.color.rgb = _rgb(fg_hex)
    run.font.size = Pt(10.5)
    run.font.name = "Calibri"


def _tinted_box(doc: Document, text: str,
                bg: str, fg: str,
                border_color: str,
                font_size: int = 10.5) -> None:
    """Single-cell table used for the flow diagram and example blocks."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _set_cell_bg(cell, bg)
    _set_cell_border(cell, border_color, size=18)

    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.font.size = Pt(font_size)
    run.font.color.rgb = _rgb(fg)
    run.font.name = "Consolas"   # monospace looks good for flows

    doc.add_paragraph()


class SalesforceTopicGenerator:
    def __init__(self, fileName: str = "Salesforce_Notes.docx"):
        print("Initializing SalesforceTopicGenerator")
        self.fileName = fileName

    # ── prompt builder ──────────────────────────────────────────────────────

    def prompt(self, prompt_type, topic,
               website_content=None, website_urls=None,
               upload_content=None):

        shared_requirements = """
            Requirements:
            - The content is intended for a Salesforce Architect.
            - Ensure all content is concise and bullet-point ready.
            - Include deep technical explanations, best practices, real-world
              architecture considerations, scalability patterns, governance
              limits, and security implications.
            - Include beginner-friendly and advanced Salesforce concepts.
            - Include Salesforce-specific terminology, object relationships,
              and API suffixes where applicable.
            - The flow_diagram must visually represent the architecture,
              workflow, hierarchy, or relationships of the topic.
            - Keep the response strictly valid JSON.
            - Do not include markdown formatting or additional explanations
              outside the JSON.
        """

        if prompt_type == "WebsiteReference":
            content_block = f"Reference Content:\n{website_content}" if website_content else ""
            urls_block    = f"Reference URLs:\n{website_urls}"       if website_urls    else ""
            return f"""
                Using the following reference material, generate a comprehensive
                Salesforce technical guide for the topic: "{topic}".
                {content_block}
                {urls_block}
                {shared_requirements}
                - Prioritize and expand on the concepts in the reference material and break down the posted content into points.
                - Supplement with your own Salesforce expertise where needed.
            """

        if prompt_type == "UploadPrompt":
            return f"""
                Using the following uploaded prompt content, generate a
                comprehensive Salesforce technical guide.
                Uploaded Content:
                {upload_content}
                {shared_requirements}
                - Follow the structure and intent of the uploaded content.
                - Supplement with Salesforce best practices where silent.
            """

        if prompt_type == "SalesforceTopic":
            return f"""
                Generate a comprehensive Salesforce technical guide for
                the topic: "{topic}".
                {shared_requirements}
                - Include examples and real-world use cases.
            """

        raise ValueError(
            f"Unknown prompt_type: '{prompt_type}'. "
            "Expected one of: 'WebsiteReference', 'UploadPrompt', 'SalesforceTopic'."
        )

    # ── document builder ────────────────────────────────────────────────────

    def generateDocument(self, json_response: dict, doc: Document):
        """
        Render json_response into a richly styled, colour-coded Word document.

        Expected keys:
            topic, summary (str or list), sections (list of {heading, content}),
            best_practices (list), limitations (list),
            flow_diagram (dict with 'representation' key),
            example (str, optional)
        """
        if json_response is None:
            print("No response — skipping document generation.")
            return

        if not doc:
            print("Document object not initialised.")
            return

        try:
            # ── 1. TITLE ──────────────────────────────────────────────────
            _banner_paragraph(
                doc,
                text=f"  {json_response.get('topic', 'Salesforce Topic')}",
                bg=COLORS["title_bg"],
                fg=COLORS["title_fg"],
                font_size=16,
            )

            # ── 2. SUMMARY ────────────────────────────────────────────────
            _banner_paragraph(doc, "  Summary",
                              COLORS["h1_bg"], COLORS["h1_fg"], font_size=12)

            summary = json_response.get("summary", "")
            if isinstance(summary, list):
                for item in summary:
                    _bullet(doc, item, COLORS["bullet_text"])
            else:
                _bullet(doc, summary, COLORS["bullet_text"])

            doc.add_paragraph()

            # ── 3. SECTIONS ───────────────────────────────────────────────
            for section in json_response.get("sections", []):
                heading = section.get("heading", "")
                content = section.get("content", "")

                _banner_paragraph(doc, f"  {heading}",
                                  COLORS["h2_bg"], COLORS["h2_fg"],
                                  font_size=11, bold=True)

                if isinstance(content, list):
                    for item in content:
                        _bullet(doc, item, COLORS["bullet_text"])
                else:
                    _bullet(doc, content, COLORS["bullet_text"])

                doc.add_paragraph()

            # ── 4. BEST PRACTICES ─────────────────────────────────────────
            _banner_paragraph(doc, "  ✔  Best Practices",
                              COLORS["h3_best"], COLORS["h3_best_fg"],
                              font_size=11)

            for practice in json_response.get("best_practices", []):
                _bullet(doc, practice, COLORS["h3_best_fg"])

            doc.add_paragraph()

            # ── 5. LIMITATIONS ────────────────────────────────────────────
            _banner_paragraph(doc, "  ⚠  Limitations",
                              COLORS["h3_limit"], COLORS["h3_limit_fg"],
                              font_size=11)

            for limitation in json_response.get("limitations", []):
                _bullet(doc, limitation, COLORS["h3_limit_fg"])

            doc.add_paragraph()

            # ── 6. DATA ACCESS / FLOW DIAGRAM ─────────────────────────────
            _banner_paragraph(doc, "  Data Access Flow",
                              COLORS["h1_bg"], COLORS["h1_fg"], font_size=12)

            flow = json_response.get("flow_diagram", {})
            representation = flow.get("representation", "")
            # Replace ASCII arrow with a nicer Unicode one
            representation = representation.replace(" -> ", "  →  ")

            _tinted_box(
                doc,
                text=representation,
                bg=COLORS["flow_bg"],
                fg="1F3864",
                border_color=COLORS["flow_border"],
                font_size=10,
            )

            # ── 7. EXAMPLE ────────────────────────────────────────────────
            example = json_response.get("example", "")
            if example:
                _banner_paragraph(doc, "  Example",
                                  COLORS["example_bg"], COLORS["example_fg"],
                                  font_size=11)
                _tinted_box(
                    doc,
                    text=example,
                    bg=COLORS["example_bg"],
                    fg=COLORS["example_fg"],
                    border_color="F4B942",
                    font_size=10,
                )

            doc.save(self.fileName)
            print(f"Document saved → {self.fileName}")
            return "Document generated successfully."

        except Exception as e:
            print(f"Error generating document: {e}")
            raise

    # ── orchestration ────────────────────────────────────────────────────
    '''
    def appendToSpecificPath(self, json_response: dict):
        #folder_path = r"D:/SF_Interview_Hub" --> used in local folder
        
        folder_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
        try:
            if not json_response:
                print("No response to append.")
                return

            full_path = os.path.join(folder_path, self.fileName)

            if not os.path.exists(folder_path):
                os.makedirs(folder_path)

            if os.path.isfile(full_path):
                print(f"Appending to existing file: {full_path}")
                doc = Document(full_path)
                # Visual page-break style separator
                doc.add_page_break()
            else:
                print(f"Creating new file: {full_path}")
                doc = Document()

            self.generateDocument(json_response, doc)
            
            return full_path 

        except Exception as e:
            print(f"Error in appendToSpecificPath: {e}")
    '''
    # Method used for render.com
    def appendToSpecificPath(self, json_response: dict):
        folder_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")

        try:
            if not json_response:
                print("No response to append.")
                return

            # Create folder FIRST before using full_path
            os.makedirs(folder_path, exist_ok=True)

            full_path = os.path.join(folder_path, self.fileName)

            if os.path.isfile(full_path):
                print(f"Appending to existing file: {full_path}")
                doc = Document(full_path)
                doc.add_page_break()
            else:
                print(f"Creating new file: {full_path}")
                doc = Document()

            self.generateDocument(json_response, doc)

            return full_path  # Return path so Flask can send the file

        except Exception as e:
            print(f"Error in appendToSpecificPath: {e}")
            raise  # Raise so Flask returns a proper 500 error
        
    def processTopic(self, prompt_type, topic,
                     website_content=None, website_urls=None,
                     upload_content=None):
        # Lazy-import so the class can be used without AIClient present
        from AIClient import AIClient as aiprocessor

        ai_client = aiprocessor("GEMENI_API_KEY")
        built_prompt = self.prompt(prompt_type, topic,
                                   website_content, website_urls,
                                   upload_content)
        print("Generated Prompt:", built_prompt)

        try:
            response = ai_client.gemeniAiConnect(prompt=built_prompt)
            print("Gemini Response:", response)
            
            if response and 'error' not in response:
                return self.appendToSpecificPath(response)
            else: 
                raise ValueError(f"Error Message : {response.get('error').get('message')}")
        except Exception as e:
            print(f"Error processing topic '{topic}': {e}")
